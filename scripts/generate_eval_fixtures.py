#!/usr/bin/env python3
"""Generate static eval fixtures under assets/eval_fixtures/."""

from __future__ import annotations

from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURE_DIR = REPO_ROOT / "assets" / "eval_fixtures"

MARKER_CSV = "EVAL_CSV_MARKER_42"
MARKER_DOCX = "EVAL_DOCX_MARKER_99"
MARKER_XLSX = "EVAL_XLSX_CELL_7"
MARKER_PDF = "EVAL_PDF_MARKER_55"
MARKER_OCR = "EVAL_OCR_MARKER"


def write_csv() -> None:
    (FIXTURE_DIR / "sample.csv").write_text(
        f"name,value\nalpha,{MARKER_CSV}\n", encoding="utf-8"
    )


def write_docx() -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph(MARKER_DOCX)
    doc.save(str(FIXTURE_DIR / "sample.docx"))


def write_xlsx() -> None:
    import pandas as pd

    df = pd.DataFrame({"col_a": [MARKER_XLSX], "col_b": ["ok"]})
    df.to_excel(FIXTURE_DIR / "sample.xlsx", index=False)


def write_pdf() -> None:
    try:
        import fitz  # PyMuPDF

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), MARKER_PDF)
        doc.save(str(FIXTURE_DIR / "sample.pdf"))
        doc.close()
    except ImportError:
        from fpdf import FPDF

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(0, 10, MARKER_PDF)
        pdf.output(str(FIXTURE_DIR / "sample.pdf"))


def write_status_eval() -> None:
    """Minimal STATUS excerpt for F4.1 workspace read eval."""
    content = """# Owlynn Status (eval fixture)

### Architectural Concerns

| Concern | Impact | Status |
|---------|--------|--------|
| Electron IPC for Screen Assist / TTS | Screen Assist and TTS require Electron main process; no browser fallback | Open — by design for desktop-only features |
| Safe Mode in browser | REST fallback via electronBridge.ts when IPC unavailable | Mitigated (BUG-5 fixed) |
| Silent error handling | Some try/catch blocks swallow errors (profile updates, API calls) | Open — partial mitigation in BUG-3/BUG-4 |
| Memory/Orchestration loading UX | Panels could hang without feedback | Mitigated (BUG-2, BUG-3 fixed — error/empty states) |
| Tool panel stale data | Mock or stale execution entries after disconnect | Mitigated (BUG-6 fixed) |
"""
    (FIXTURE_DIR / "status_eval.md").write_text(content, encoding="utf-8")


def write_ocr_png() -> None:
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (400, 120), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 28)
    except OSError:
        font = ImageFont.load_default()
    draw.text((20, 40), MARKER_OCR, fill=(0, 0, 0), font=font)
    img.save(FIXTURE_DIR / "ocr_sample.png")


def main() -> None:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    write_csv()
    write_docx()
    write_xlsx()
    write_pdf()
    write_status_eval()
    write_ocr_png()
    print(f"Fixtures written to {FIXTURE_DIR}")


if __name__ == "__main__":
    main()
